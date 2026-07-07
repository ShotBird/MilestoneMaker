#!/usr/bin/env python3
"""Schedule-deck engine (vendored from generate_schedule_pptx.mjs/.py).

Reads a schedule xlsx (openpyxl) and edits the reference template deck
(python-pptx) to produce an editable schedule .pptx -- no @oai/artifact-tool.

All internal geometry uses the template's 1280-unit coordinate system, the same
constants as the .mjs. python-pptx works in EMU, so 1 unit = 9525 EMU
(1 unit = 0.75 pt, 1 pt = 12700 EMU).

Public API: ``generate_schedule(input_path, template_path, output_path)`` does
the read -> edit -> save and returns a summary dict. ``main()`` is a thin CLI
wrapper around it so the engine stays runnable standalone (used by tests).
The geometry constants and logic are unchanged from the verified port.
"""

import argparse
import copy
import datetime as dt
import json
import os
import re
import sys
from pathlib import Path

import openpyxl
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

UNIT = 9525.0  # EMU per 1280-unit


# --------------------------------------------------------------------------- #
# Bundled-resource resolution (works both frozen by PyInstaller and from source)
# --------------------------------------------------------------------------- #
def resource_path(rel):
    """Absolute path to a bundled resource (e.g. the embedded template)."""
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, rel)


def app_dir():
    """Directory to write side-files (error.log) next to the exe / source."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


# --------------------------------------------------------------------------- #
# Date helpers (Excel 1900 system, epoch 1899-12-30, matching the .mjs)
# --------------------------------------------------------------------------- #
EPOCH = dt.date(1899, 12, 30)


def excel_serial_to_date(serial):
    if not isinstance(serial, (int, float)) or isinstance(serial, bool):
        return None
    return EPOCH + dt.timedelta(days=int(serial))


def date_to_excel_serial(d):
    return (dt.date(d.year, d.month, d.day) - EPOCH).days


def month_start_serial(serial):
    d = excel_serial_to_date(serial)
    return date_to_excel_serial(dt.date(d.year, d.month, 1)) if d else serial


def month_end_serial(serial):
    d = excel_serial_to_date(serial)
    if not d:
        return serial
    nm = dt.date(d.year + (d.month // 12), (d.month % 12) + 1, 1)
    return date_to_excel_serial(nm - dt.timedelta(days=1))


def format_date(serial):
    d = excel_serial_to_date(serial)
    return d.isoformat() if d else str(serial)


def day_of_month(serial):
    d = excel_serial_to_date(serial)
    return str(d.day) if d else ""


def year_label(serial):
    d = excel_serial_to_date(serial)
    return f"{str(d.year)[2:]}년" if d else ""


def year_range_label(start_serial, end_serial):
    s = excel_serial_to_date(start_serial)
    e = excel_serial_to_date(end_serial)
    if not s or not e:
        return ""
    sy, ey = str(s.year)[2:], str(e.year)[2:]
    return f"{sy}년" if sy == ey else f"{sy}-{ey}년"


def month_label(serial, base_year):
    d = excel_serial_to_date(serial)
    if not d:
        return ""
    if d.year == base_year:
        return f"{d.month}월"
    return f"{str(d.year)[2:]}.{d.month}월"


def month_key(serial):
    d = excel_serial_to_date(serial)
    return d.year * 12 + (d.month - 1) if d else None


def serial_from_month_key(key):
    year = key // 12
    month = key % 12
    return date_to_excel_serial(dt.date(year, month + 1, 1))


def clamp(value, lo, hi):
    return min(max(value, lo), hi)


# --------------------------------------------------------------------------- #
# Workbook readers
# --------------------------------------------------------------------------- #
def read_values(ws):
    """Return A1:Q16 as values[row][col] (0-based), dates -> Excel serials."""
    values = []
    for r in range(1, 17):
        row = []
        for c in range(1, 18):  # columns A..Q
            v = ws.cell(row=r, column=c).value
            if isinstance(v, dt.datetime):
                v = date_to_excel_serial(v.date())
            elif isinstance(v, dt.date):
                v = date_to_excel_serial(v)
            row.append(v)
        values.append(row)
    return values


def _cell(values, row, col):
    if 0 <= row < len(values) and 0 <= col < len(values[row]):
        return values[row][col]
    return None


def is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def read_items(values, rows, label_col, date_col):
    out = []
    for row in rows:
        label = _cell(values, row, label_col)
        date = _cell(values, row, date_col)
        if label not in (None, "") and is_number(date):
            out.append({"label": label, "date": date})
    return out


def readable_cell_value(v):
    if v is None:
        return ""
    # normalise in-cell line breaks (Alt+Enter) to LF; a stray CR otherwise renders
    # as an extra blank line and inflates label-width estimates by one glyph.
    return str(v).replace("\r\n", "\n").replace("\r", "\n").strip()


def read_subcolumn_lane(values, rows, group_label_col, target_label_col,
                        target_date_col, check_label_col, check_date_col):
    labels = []
    groups = []
    state = {"current": None}

    def ensure_group(label=""):
        if state["current"] is None or label:
            state["current"] = {"label": label, "targets": [], "events": [], "_seg": 0, "_brk": False}
            groups.append(state["current"])
            if label and label not in labels:
                labels.append(label)
        return state["current"]

    for row in rows:
        group_label = readable_cell_value(_cell(values, row, group_label_col))
        if group_label:
            ensure_group(group_label)

        target_label = readable_cell_value(_cell(values, row, target_label_col))
        target_date = _cell(values, row, target_date_col)
        has_target = bool(target_label) and is_number(target_date)
        if has_target:
            g = ensure_group()
            g["targets"].append({"label": target_label, "date": target_date,
                                  "group": g["label"], "source": "target"})

        check_label = readable_cell_value(_cell(values, row, check_label_col))
        check_date = _cell(values, row, check_date_col)
        has_check = bool(check_label) and is_number(check_date)
        if has_check:
            g = ensure_group()
            if g["_brk"]:                       # a blank row was seen -> start a new arrow segment
                g["_seg"] += 1
                g["_brk"] = False
            g["events"].append({"label": check_label, "date": check_date,
                                "group": g["label"], "source": "check", "seg": g["_seg"]})

        # a fully blank row inside a group that already has check-posts -> the next check-posts
        # begin a NEW segment (a separate connector arrow): split-evaluation support.
        if (not group_label and not has_target and not has_check
                and state["current"] is not None and state["current"]["events"]):
            state["current"]["_brk"] = True

    events = []
    out_groups = []
    for g in groups:
        group_events = list(g["events"])
        last_seg = max((e["seg"] for e in g["events"]), default=0)
        for target in g["targets"]:
            present = any(e["label"] == target["label"] and e["date"] == target["date"]
                          for e in group_events)
            if not present:
                t = dict(target)
                t["source"] = "target-as-final-check"
                t["seg"] = last_seg
                group_events.append(t)
        events.extend(group_events)
        # one entry PER group-label row -> two evaluation items with the same name stay separate
        # (and keep their own arrow segments) instead of being merged by name.
        out_groups.append({"name": g["label"], "events": group_events})
    return {"labels": labels, "events": events, "groups": out_groups}


# --------------------------------------------------------------------------- #
# Template object map (artifact-tool ids -> python-pptx shape names)
# --------------------------------------------------------------------------- #
TEMPLATE_NAMES = {
    "title": "Text Box 36",
    "punchline": "Content Placeholder 8",
    "table": "표 4",
    "now_marker": "이등변 삼각형 18",
    "now_label": "TextBox 19",
    "gen_label": "직사각형 32",
    "gen_circle": "타원 33",
    # arrows share a name ("오른쪽 화살표 96"); disambiguated by top.
    "product": [
        ("직사각형 16", "타원 17"),
        ("직사각형 26", "타원 27"),
        ("직사각형 76", "타원 77"),
        ("직사각형 13", "타원 14"),
        ("직사각형 7", "타원 8"),
    ],
    "lane1": [
        ("직사각형 32", "타원 33"),
        ("직사각형 34", "타원 35"),
        ("직사각형 36", "타원 37"),
        ("직사각형 24", "타원 25"),
        ("직사각형 28", "타원 29"),
        ("직사각형 30", "타원 31"),
        ("직사각형 38", "타원 39"),
    ],
    "lane2": [
        ("직사각형 40", "타원 41"),
        ("직사각형 42", "타원 43"),
    ],
}
ARROW_NAME = "오른쪽 화살표 96"


# --------------------------------------------------------------------------- #
# Shape wrapper -- exposes position in 1280-units, like the .mjs shape.position
# --------------------------------------------------------------------------- #
class S:
    def __init__(self, sp):
        self.sp = sp

    @property
    def left(self):
        return (self.sp.left or 0) / UNIT

    @left.setter
    def left(self, v):
        self.sp.left = Emu(int(round(v * UNIT)))

    @property
    def top(self):
        return (self.sp.top or 0) / UNIT

    @top.setter
    def top(self, v):
        self.sp.top = Emu(int(round(v * UNIT)))

    @property
    def width(self):
        return (self.sp.width or 0) / UNIT

    @width.setter
    def width(self, v):
        self.sp.width = Emu(int(round(v * UNIT)))

    @property
    def height(self):
        return (self.sp.height or 0) / UNIT

    @height.setter
    def height(self, v):
        self.sp.height = Emu(int(round(v * UNIT)))

    def set_text(self, text, color=None, bold=None, size_pt=None):
        """Set text, preserving the first run's font unless overrides are given."""
        tf = self.sp.text_frame
        para = tf.paragraphs[0]
        # remove extra paragraphs
        for extra in list(tf.paragraphs[1:]):
            extra._p.getparent().remove(extra._p)
        runs = para.runs
        if runs:
            run = runs[0]
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
        else:
            run = para.add_run()
        run.text = str(text)
        if color is not None:
            run.font.color.rgb = color
        if bold is not None:
            run.font.bold = bold
        if size_pt is not None:
            run.font.size = Pt(size_pt)


class Resolver:
    def __init__(self, slide):
        self.slide = slide
        self.by_name = {}
        for sp in slide.shapes:
            self.by_name.setdefault(sp.name, []).append(sp)

    def by_name_one(self, name):
        items = self.by_name.get(name)
        if not items:
            raise KeyError(f"template shape not found: {name!r}")
        return S(items[0])

    def arrow(self, kind):
        arrows = sorted(self.by_name.get(ARROW_NAME, []), key=lambda s: s.top or 0)
        if len(arrows) < 2:
            raise KeyError("expected two arrows named " + ARROW_NAME)
        return S(arrows[0] if kind == "lane1" else arrows[1])  # lane1 is the upper arrow (smaller top)


# --------------------------------------------------------------------------- #
# Calibration constants (identical to the .mjs)
# --------------------------------------------------------------------------- #
TIMELINE_ANCHORS = sorted([
    [46037, 224.5], [46110, 473.0], [46117, 532.9], [46120, 574.8],
    [46127, 613.0], [46145, 722.3], [46151, 767.0], [46157, 805.6],
    [46188, 952.0], [46203, 1018.8], [46264, 1254.5],
], key=lambda a: a[0])

TABLE_GEOMETRY = {
    "left": 18.41,
    "column_widths": [65.41, 65.41, 135.02, 68.75, 153.0, 208.0, 186.0, 150.0, 77.62, 140.56],
    "month_start_column": 2,
    "month_column_count": 8,
}
CIRCLE_PADDING = 15.12

REFERENCE_ROWS = {
    "product_height": 76.0142782152231,
    "expanded_product_height": 92,
    # 직접 입력 row matches the template marker-lane height; 2-line marker labels
    # overflow downward exactly like the template marker lanes do.
    "direct_row_height": 120,
    "fitted_direct_row_height": 59.12,
}

GENERATED_LANE_PLACEMENTS = {
    # direct tops use the same offsets the first marker lane has from its row top (~1.6 / ~32.8).
    "direct": {"circle_top": 358.64, "arrow_top": 360.53, "label_top": 389.84, "max_items": 8},
    "lane3": {"circle_top": 504.9, "arrow_top": 506.79, "label_top": 538.0, "max_items": 8},
}


def timeline_frame_left():
    g = TABLE_GEOMETRY
    return g["left"] + sum(g["column_widths"][:g["month_start_column"]])


def timeline_frame_right():
    g = TABLE_GEOMETRY
    return g["left"] + sum(g["column_widths"])


def timeline_month_columns():
    g = TABLE_GEOMETRY
    left = timeline_frame_left()
    cols = []
    for index, width in enumerate(
        g["column_widths"][g["month_start_column"]:g["month_start_column"] + g["month_column_count"]]
    ):
        cols.append({"cell_index": g["month_start_column"] + index, "left": left, "width": width})
        left += width
    return cols


# --------------------------------------------------------------------------- #
# Timeline scale
# --------------------------------------------------------------------------- #
def create_timeline_scale(dates, warnings):
    if not dates:
        return {"type": "anchored"}
    min_date, max_date = min(dates), max(dates)
    first_anchor = TIMELINE_ANCHORS[0][0]
    last_anchor = TIMELINE_ANCHORS[-1][0]
    if min_date >= first_anchor and max_date <= last_anchor:
        return {"type": "anchored"}
    start_serial = month_start_serial(min_date)
    end_serial = month_end_serial(max_date)
    warnings.append(
        f"Dates extend outside the calibrated reference timeline; using fitted timeline "
        f"{format_date(start_serial)} to {format_date(end_serial)}."
    )
    return {"type": "fitted", "start_serial": start_serial, "end_serial": end_serial}


def month_keys_between(start_serial, end_serial):
    sk, ek = month_key(start_serial), month_key(end_serial)
    if sk is None or ek is None:
        return []
    return list(range(sk, ek + 1))


def select_header_month_keys(scale, dates):
    column_count = TABLE_GEOMETRY["month_column_count"]
    all_keys = month_keys_between(scale["start_serial"], scale["end_serial"])
    counts = {}
    for date in dates:
        k = month_key(date)
        if k is not None:
            counts[k] = counts.get(k, 0) + 1

    event_keys = sorted(counts.keys())
    if len(event_keys) >= column_count:
        selected = {event_keys[0], event_keys[-1]}
        middle = sorted(event_keys[1:-1], key=lambda k: (-counts[k], k))
        for k in middle[:max(0, column_count - len(selected))]:
            selected.add(k)
        return sorted(selected)[:column_count]

    selected = set(event_keys)
    missing = [k for k in all_keys if k not in selected]
    needed = column_count - len(selected)
    for index in range(0, max(0, needed)):
        if not missing:
            break
        missing_index = round(((index + 1) * (len(missing) - 1)) / (needed + 1))
        selected.add(missing[missing_index])
    return sorted(selected)[:column_count]


def selected_month_keys_for_scale(scale):
    if scale.get("header_month_keys"):
        return scale["header_month_keys"]
    return month_keys_between(scale["start_serial"], scale["end_serial"])[:TABLE_GEOMETRY["month_column_count"]]


def x_within_month_column(serial, column):
    key = month_key(serial)
    start_serial = serial_from_month_key(key)
    end_serial = month_end_serial(start_serial)
    left = column["left"] + CIRCLE_PADDING
    right = column["left"] + column["width"] - CIRCLE_PADDING
    ratio = clamp((serial - start_serial) / max(1, end_serial - start_serial), 0, 1)
    return left + ratio * max(0, right - left)


def x_for_fitted_date_raw(serial, scale):
    columns = timeline_month_columns()
    selected_keys = selected_month_keys_for_scale(scale)
    key = month_key(serial)
    if key in selected_keys:
        idx = selected_keys.index(key)
        if idx < len(columns):
            return x_within_month_column(serial, columns[idx])

    points = []
    for index, sel in enumerate(selected_keys):
        if index < len(columns):
            points.append({"serial": serial_from_month_key(sel),
                           "x": columns[index]["left"] + columns[index]["width"] / 2})
    if not points:
        ratio = (serial - scale["start_serial"]) / max(1, scale["end_serial"] - scale["start_serial"])
        return timeline_frame_left() + ratio * (timeline_frame_right() - timeline_frame_left())

    if serial <= points[0]["serial"]:
        return points[0]["x"]
    last = points[-1]
    if serial >= last["serial"]:
        return last["x"]
    for index in range(len(points) - 1):
        lp, rp = points[index], points[index + 1]
        if lp["serial"] <= serial <= rp["serial"]:
            ratio = (serial - lp["serial"]) / max(1, rp["serial"] - lp["serial"])
            return lp["x"] + ratio * (rp["x"] - lp["x"])
    return last["x"]


def spread_sorted_centers_within_bounds(centers, minimum_gap, lo, hi):
    if len(centers) < 2:
        return [clamp(c, lo, hi) for c in centers]
    feasible_gap = max(0, (hi - lo) / max(1, len(centers) - 1))
    gap = min(minimum_gap, feasible_gap)
    adjusted = [clamp(c, lo, hi) for c in centers]
    for i in range(1, len(adjusted)):
        adjusted[i] = max(adjusted[i], adjusted[i - 1] + gap)
    overflow = adjusted[-1] - hi
    if overflow > 0:
        adjusted = [a - overflow for a in adjusted]
    for i in range(len(adjusted) - 2, -1, -1):
        adjusted[i] = min(adjusted[i], adjusted[i + 1] - gap)
    underflow = lo - adjusted[0]
    if underflow > 0:
        adjusted = [a + underflow for a in adjusted]
    return [clamp(c, lo, hi) for c in adjusted]


def build_date_center_overrides(dates, scale):
    if scale["type"] != "fitted":
        return {"centers": {}, "locked_serials": set()}
    columns = timeline_month_columns()
    selected_keys = selected_month_keys_for_scale(scale)
    date_counts = {}
    for serial in dates:
        if is_number(serial):
            date_counts[serial] = date_counts.get(serial, 0) + 1
    unique_dates = sorted(date_counts.keys())

    dates_by_month = {}
    for serial in unique_dates:
        key = month_key(serial)
        if key not in selected_keys:
            continue
        column_index = selected_keys.index(key)
        if column_index >= len(columns):
            continue
        column = columns[column_index]
        dates_by_month.setdefault(key, {"column": column, "serials": []})["serials"].append(serial)

    overrides = {}
    locked_serials = {s for s, c in date_counts.items() if c > 1}
    minimum_gap = 66
    for entry in dates_by_month.values():
        column = entry["column"]
        lo = column["left"] + CIRCLE_PADDING
        hi = column["left"] + column["width"] - CIRCLE_PADDING
        centers = [x_for_fitted_date_raw(s, scale) for s in entry["serials"]]
        adjusted = spread_sorted_centers_within_bounds(centers, minimum_gap, lo, hi)
        for serial, center in zip(entry["serials"], adjusted):
            overrides[serial] = center
    return {"centers": overrides, "locked_serials": locked_serials}


def x_for_fitted_date(serial, scale):
    ov = scale.get("date_center_overrides")
    if ov and ov.get("centers"):
        centers = ov["centers"]
        if serial in centers:
            return centers[serial]
        # A non-milestone date (e.g. the Now marker) must stay consistent with the
        # spread circle positions: interpolate between the nearest milestone override
        # centres by serial (those are monotonic), so 'today' never lands left of an
        # earlier milestone that was pushed right by label spacing.
        keys = sorted(centers)
        lo = hi = None
        for k in keys:
            if k <= serial:
                lo = k
            if k >= serial and hi is None:
                hi = k
        if lo is not None and hi is not None:
            if lo == hi:
                return centers[lo]
            ratio = (serial - lo) / (hi - lo)
            return centers[lo] + ratio * (centers[hi] - centers[lo])
        if lo is not None:
            return centers[lo]
        if hi is not None:
            return centers[hi]
    return x_for_fitted_date_raw(serial, scale)


def should_clamp_label_to_month(label):
    return bool(re.match(r"^[A-Za-z]+/[A-Za-z]+$", str(label).strip()))


def month_column_bounds_for_date(serial, scale, label):
    if not should_clamp_label_to_month(label):
        return None
    if scale["type"] != "fitted":
        return None
    columns = timeline_month_columns()
    selected_keys = selected_month_keys_for_scale(scale)
    key = month_key(serial)
    if key not in selected_keys:
        return None
    idx = selected_keys.index(key)
    if idx >= len(columns):
        return None
    column = columns[idx]
    return {"left": column["left"] + 2, "right": column["left"] + column["width"] - 2}


def x_for_date(serial, scale, warnings, warn=True):
    if scale["type"] == "fitted":
        x = x_for_fitted_date(serial, scale)
        return clamp(x, timeline_frame_left() + CIRCLE_PADDING, timeline_frame_right() - CIRCLE_PADDING)

    if serial < TIMELINE_ANCHORS[0][0]:
        if warn:
            warnings.append(f"Date {serial} is before the calibrated timeline; clamped to the first anchor.")
        return TIMELINE_ANCHORS[0][1]
    last = TIMELINE_ANCHORS[-1]
    if serial > last[0]:
        if warn:
            warnings.append(f"Date {serial} is after the calibrated timeline; clamped to the last anchor.")
        return last[1]
    for index in range(len(TIMELINE_ANCHORS) - 1):
        ls, lx = TIMELINE_ANCHORS[index]
        rs, rx = TIMELINE_ANCHORS[index + 1]
        if ls <= serial <= rs:
            ratio = (serial - ls) / (rs - ls)
            return lx + ratio * (rx - lx)
    return last[1]


def spread_dense_centers(centers, lane_name, scale, warnings):
    if scale["type"] != "fitted" or scale.get("date_center_overrides") or len(centers) < 2:
        return centers
    lo = timeline_frame_left() + CIRCLE_PADDING
    hi = timeline_frame_right() - CIRCLE_PADDING
    lane_min_gaps = {"product": 42, "lane1": 66, "lane2": 56, "direct": 66, "lane3": 66}
    minimum_gap = min(lane_min_gaps.get(lane_name, 42), (hi - lo) / max(1, len(centers) - 1))
    entries = sorted([{"center": c, "index": i} for i, c in enumerate(centers)],
                     key=lambda e: (e["center"], e["index"]))
    adjusted = [e["center"] for e in entries]
    for i in range(1, len(adjusted)):
        adjusted[i] = max(adjusted[i], adjusted[i - 1] + minimum_gap)
    overflow = adjusted[-1] - hi
    if overflow > 0:
        adjusted = [a - overflow for a in adjusted]
    for i in range(len(adjusted) - 2, -1, -1):
        adjusted[i] = min(adjusted[i], adjusted[i + 1] - minimum_gap)
    underflow = lo - adjusted[0]
    if underflow > 0:
        adjusted = [a + underflow for a in adjusted]
    by_original = list(centers)
    moved = False
    for sorted_index, entry in enumerate(entries):
        next_center = clamp(adjusted[sorted_index], lo, hi)
        if abs(next_center - centers[entry["index"]]) > 0.5:
            moved = True
        by_original[entry["index"]] = next_center
    if moved:
        warnings.append(f"{lane_name} markers were spaced slightly to avoid overlap on the fitted timeline.")
    return by_original


def centers_for_items(items, lane_name, scale, warnings):
    centers = [x_for_date(item["date"], scale, warnings) for item in items]
    return spread_dense_centers(centers, lane_name, scale, warnings)


def fitted_label_width(text, current_width, scale, lane=None):
    if lane == "product":
        # Keep the short/medium codes (3-char codes stay ~54 -> no over-widening, no lane shift);
        # only widen LONG product names (e.g. a parenthesised Korean name) so they stay on
        # one line. CJK chars count wider than Latin.
        if len(text) <= 3:
            return min(max(current_width, 46), 54)
        if len(text) <= 5:
            return max(current_width, 72)
        cjk = len(re.findall(r"[　-鿿가-힯豈-﫿]", text))
        estimated = cjk * 17 + (len(text) - cjk) * 9 + 20
        return max(current_width, estimated)
    if scale["type"] != "fitted":
        return current_width
    longest_line = max(len(line) for line in re.split(r"\r?\n", text))
    if "\n" in text:
        return max(current_width, longest_line * 9 + 30, 110)
    if len(text) <= 3:
        return min(max(current_width, 46), 54)
    if len(text) <= 5:
        return max(current_width, 72)
    if re.search(r"cyc$", text, re.IGNORECASE):
        return max(current_width, len(text) * 8 + 18)
    return current_width


def fitted_label_height(text, current_height):
    return max(current_height, 43.23) if "\n" in text else current_height


# --------------------------------------------------------------------------- #
# Placement
# --------------------------------------------------------------------------- #
def place_pair_at_center(label, circle, item, circle_center, scale, placement=None):
    placement = placement or {}
    label_text = str(item["label"])
    label.set_text(label_text)
    circle.set_text(day_of_month(item["date"]), color=RGBColor(0xFF, 0xFF, 0xFF),
                    bold=True, size_pt=14.67)

    label_width = fitted_label_width(label_text, label.width, scale, placement.get("lane"))
    label_height = fitted_label_height(label_text, label.height)
    bounds = placement.get("bounds")
    if bounds and label_width <= bounds["right"] - bounds["left"]:
        use_bounds = bounds
    else:
        use_bounds = {"left": 0, "right": 1280}
    label_left = clamp(circle_center - label_width / 2, use_bounds["left"], use_bounds["right"] - label_width)
    adjusted_center = circle_center if placement.get("lock_circle_center") else label_left + label_width / 2
    circle.left = adjusted_center - circle.width / 2
    if placement.get("circle_top") is not None:
        circle.top = placement["circle_top"]
    label.left = label_left
    if placement.get("label_top") is not None:
        label.top = placement["label_top"]
    label.width = label_width
    label.height = label_height
    return adjusted_center


def deconflict_lane_labels(placed, scale):
    if scale["type"] != "fitted" or len(placed) < 2:
        return
    gap = 4
    lo, hi = 0, 1280
    entries = sorted(
        [{"label": p["label"], "index": i, "left": p["label"].left, "width": p["label"].width}
         for i, p in enumerate(placed)],
        key=lambda e: (e["left"], e["index"]),
    )
    for i in range(1, len(entries)):
        entries[i]["left"] = max(entries[i]["left"], entries[i - 1]["left"] + entries[i - 1]["width"] + gap)
    overflow = entries[-1]["left"] + entries[-1]["width"] - hi
    if overflow > 0:
        for e in entries:
            e["left"] -= overflow
    for i in range(len(entries) - 2, -1, -1):
        entries[i]["left"] = min(entries[i]["left"], entries[i + 1]["left"] - gap - entries[i]["width"])
    underflow = lo - entries[0]["left"]
    if underflow > 0:
        for e in entries:
            e["left"] += underflow
    for e in entries:
        e["label"].left = clamp(e["left"], lo, hi - e["width"])


def align_unlocked_circles_to_labels(placed, scale):
    final = []
    locked_set = scale.get("date_center_overrides", {}).get("locked_serials", set()) if scale.get("date_center_overrides") else set()
    for p in placed:
        label, circle, item = p["label"], p["circle"], p["item"]
        label_center = label.left + label.width / 2
        locked = item["date"] in locked_set
        if circle is not None and not locked:
            circle.left = label_center - circle.width / 2
        final.append(circle.left + circle.width / 2 if circle is not None else label_center)
    return final


def move_lane_arrow(arrow, centers, mode):
    if len(centers) < 2:
        return
    left_center, right_center = min(centers), max(centers)
    delta = right_center - left_center
    if mode == "lane1":
        arrow.left = left_center + 16.5
        arrow.width = max(48, delta - 31.5)
        return
    arrow.left = left_center - 7.2
    arrow.width = max(40, delta + 22.3)


def update_lane(resolver, lane_name, items, scale, warnings):
    pairs = TEMPLATE_NAMES[lane_name]
    if len(items) > len(pairs):
        warnings.append(f"{lane_name} has {len(items)} items but the template supports "
                        f"{len(pairs)}; extra items were skipped.")
    supported = items[:len(pairs)]
    centers = centers_for_items(supported, lane_name, scale, warnings)
    locked_set = scale.get("date_center_overrides", {}).get("locked_serials", set()) if scale.get("date_center_overrides") else set()
    placed = []
    for index, (label_name, circle_name) in enumerate(pairs):
        if index >= len(supported):
            break
        item = supported[index]
        label = resolver.by_name_one(label_name)
        circle = resolver.by_name_one(circle_name)
        place_pair_at_center(label, circle, item, centers[index], scale, {
            "bounds": month_column_bounds_for_date(item["date"], scale, item["label"]),
            "lock_circle_center": item["date"] in locked_set,
            "lane": lane_name,
        })
        placed.append({"label": label, "circle": circle, "item": item})
    deconflict_lane_labels(placed, scale)
    final_centers = align_unlocked_circles_to_labels(placed, scale)
    if lane_name in ("lane1", "lane2"):
        move_lane_arrow(resolver.arrow(lane_name), final_centers, lane_name)


def should_expand_product_row(items):
    for item in items:
        label = str(item["label"])
        if len(label) > 10 or re.search(r"[()]", label):
            return True
    return False


def update_product_row_height(table, items):
    if not should_expand_product_row(items):
        return 0
    set_row_height(table, 1, REFERENCE_ROWS["expanded_product_height"])
    return REFERENCE_ROWS["expanded_product_height"] - REFERENCE_ROWS["product_height"]


def fit_direct_row(table):
    set_row_height(table, 4, REFERENCE_ROWS["fitted_direct_row_height"])
    return REFERENCE_ROWS["direct_row_height"] - REFERENCE_ROWS["fitted_direct_row_height"]


def set_row_height(table, index, value_1280):
    table.rows[index].height = Emu(int(round(value_1280 * UNIT)))


def shift_template_lane_vertical(resolver, lane_name, delta):
    if not delta:
        return
    for label_name, circle_name in TEMPLATE_NAMES[lane_name]:
        for name in (label_name, circle_name):
            sh = resolver.by_name_one(name)
            sh.top = sh.top + delta
    arrow = resolver.arrow(lane_name)
    arrow.top = arrow.top + delta


# --------------------------------------------------------------------------- #
# Shape cloning for generated lanes
# --------------------------------------------------------------------------- #
def _next_shape_id(slide):
    ids = [0]
    for el in slide.shapes._spTree.iter(qn("p:cNvPr")):
        try:
            ids.append(int(el.get("id")))
        except (TypeError, ValueError):
            pass
    return max(ids) + 1


def clone_shape(slide, source, name, top):
    new_el = copy.deepcopy(source.sp._element)
    cnvpr = new_el.find(".//" + qn("p:cNvPr"))
    cnvpr.set("id", str(_next_shape_id(slide)))
    cnvpr.set("name", name)
    slide.shapes._spTree.append(new_el)
    # locate the new python-pptx shape by name
    for sp in slide.shapes:
        if sp.name == name:
            wrapper = S(sp)
            wrapper.top = top
            return wrapper
    raise RuntimeError("clone failed: " + name)


def add_generated_lane(slide, resolver, lane_name, items, scale, warnings, vertical_offset=0):
    placement = GENERATED_LANE_PLACEMENTS.get(lane_name)
    if not placement or not items:
        return
    if len(items) > placement["max_items"]:
        warnings.append(f"{lane_name} has {len(items)} items but the generated lane supports "
                        f"{placement['max_items']}; extra items were skipped.")
    supported = items[:placement["max_items"]]
    centers = centers_for_items(supported, lane_name, scale, warnings)
    locked_set = scale.get("date_center_overrides", {}).get("locked_serials", set()) if scale.get("date_center_overrides") else set()

    arrow = None
    if len(centers) > 1:
        arrow = clone_shape(slide, resolver.arrow("lane1"), f"{lane_name} generated arrow",
                            placement["arrow_top"] + vertical_offset)

    placed = []
    gen_circle_src = resolver.by_name_one(TEMPLATE_NAMES["gen_circle"])
    gen_label_src = resolver.by_name_one(TEMPLATE_NAMES["gen_label"])
    for index, item in enumerate(supported):
        circle = clone_shape(slide, gen_circle_src, f"{lane_name} generated circle {index + 1}",
                             placement["circle_top"] + vertical_offset)
        label = clone_shape(slide, gen_label_src, f"{lane_name} generated label {index + 1}",
                            placement["label_top"] + vertical_offset)
        place_pair_at_center(label, circle, item, centers[index], scale, {
            "bounds": month_column_bounds_for_date(item["date"], scale, item["label"]),
            "circle_top": placement["circle_top"] + vertical_offset,
            "label_top": placement["label_top"] + vertical_offset,
            "lock_circle_center": item["date"] in locked_set,
            "lane": lane_name,
        })
        placed.append({"label": label, "circle": circle, "item": item})
    deconflict_lane_labels(placed, scale)
    final_centers = align_unlocked_circles_to_labels(placed, scale)
    if arrow is not None:
        move_lane_arrow(arrow, final_centers, "lane1")


# --------------------------------------------------------------------------- #
# Table text
# --------------------------------------------------------------------------- #
def set_cell_text(cell, text, bold=None, size_pt=None):
    """Set cell text (supports \\n lines), preserving the first run's font.
    Normalise CRLF/CR -> LF so a stray carriage return never renders as an extra
    blank line, and drop leading/trailing blank lines (e.g. a trailing Alt+Enter)."""
    s = str(text).replace("\r\n", "\n").replace("\r", "\n")
    lines = s.split("\n")
    while lines and lines[0].strip() == "":
        lines.pop(0)
    while lines and lines[-1].strip() == "":
        lines.pop()
    if not lines:
        lines = [""]
    tf = cell.text_frame
    # keep only the first paragraph
    for extra in list(tf.paragraphs[1:]):
        extra._p.getparent().remove(extra._p)
    first = tf.paragraphs[0]
    runs = first.runs
    if runs:
        run = runs[0]
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        run = first.add_run()
    run.text = lines[0]
    template_run = run
    for line in lines[1:]:
        para = tf.add_paragraph()
        new_run = para.add_run()
        new_run.text = line
        _copy_run_font(template_run, new_run)
    for para in tf.paragraphs:
        for r in para.runs:
            if bold is not None:
                r.font.bold = bold
            if size_pt is not None:
                r.font.size = Pt(size_pt)


def _copy_run_font(src_run, dst_run):
    src_rpr = src_run._r.find(qn("a:rPr"))
    if src_rpr is None:
        return
    dst_r = dst_run._r
    old = dst_r.find(qn("a:rPr"))
    if old is not None:
        dst_r.remove(old)
    dst_r.insert(0, copy.deepcopy(src_rpr))


def _subcolumn_size(labels):
    # artifact-tool sets these to 16 and auto-shrinks to fit the narrow cell; python-pptx
    # has no autofit, so bake the shrink (long 8-char names -> 12) and keep short
    # names at 16, matching the reference deck.
    longest = max((len(l) for l in labels), default=0)
    return 12 if longest > 6 else 16


def update_subcolumn_labels(table, data):
    set_cell_text(table.cell(4, 1), "\n".join(data["direct_labels"]), bold=True, size_pt=_subcolumn_size(data["direct_labels"]))
    set_cell_text(table.cell(5, 1), "\n".join(data["pra_labels"]), bold=True, size_pt=_subcolumn_size(data["pra_labels"]))


def update_timeline_header(table, scale, dates):
    if not dates:
        return
    if scale["type"] != "fitted":
        set_cell_text(table.cell(0, 0), year_label(min(dates)))
        return
    set_cell_text(table.cell(0, 0), year_range_label(scale["start_serial"], scale["end_serial"]))
    base_year = excel_serial_to_date(scale["start_serial"]).year
    selected = selected_month_keys_for_scale(scale)
    for index, column in enumerate(timeline_month_columns()):
        key = selected[index] if index < len(selected) else None
        text = "" if key is None else month_label(serial_from_month_key(key), base_year)
        set_cell_text(table.cell(0, column["cell_index"]), text)


def today_serial():
    now = dt.date.today()
    return date_to_excel_serial(now)


def move_now_marker(resolver, scale):
    if scale["type"] != "fitted":
        return
    serial = today_serial()
    marker = resolver.by_name_one(TEMPLATE_NAMES["now_marker"])
    label = resolver.by_name_one(TEMPLATE_NAMES["now_label"])
    center = x_for_date(serial, scale, [], warn=False)
    marker.left = center - marker.width / 2
    label.left = center + 12.1


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #
def get_table(resolver):
    return resolver.by_name_one(TEMPLATE_NAMES["table"]).sp.table


def generate_schedule(input_path, template_path, output_path):
    """Read an xlsx, edit the template deck, save ``output_path``.

    Returns a summary dict ``{"output", "warnings", "counts"}``. Raises on bad
    input (the GUI caller turns the exception into a dialog + error.log). The
    body below is the verified port -- geometry/logic unchanged.
    """
    warnings = []

    wb = openpyxl.load_workbook(input_path, data_only=True)
    ws = wb.worksheets[0]
    values = read_values(ws)
    direct_lane = read_subcolumn_lane(values, [5, 6, 7, 8, 9, 10, 11], 7, 8, 9, 10, 11)
    pra_lane = read_subcolumn_lane(values, [5, 6, 7, 8, 9, 10, 11], 12, 13, 14, 15, 16)

    data = {
        "title": _cell(values, 0, 1) or _cell(values, 0, 0) or "",
        "punchline": _cell(values, 1, 1) or "",
        "product": read_items(values, [4, 5, 6, 7, 8], 0, 1),
        "lane1": read_items(values, [5, 6, 7, 8, 9, 10, 11], 5, 6),
        "lane2": read_items(values, [14, 15], 5, 6),
        "direct_labels": direct_lane["labels"],
        "direct": direct_lane["events"],
        "pra_labels": pra_lane["labels"],
        "lane3": pra_lane["events"],
    }

    prs = Presentation(template_path)
    slide = prs.slides[0]
    resolver = Resolver(slide)
    table = get_table(resolver)

    resolver.by_name_one(TEMPLATE_NAMES["title"]).set_text(f" {data['title']}")
    resolver.by_name_one(TEMPLATE_NAMES["punchline"]).set_text(f" ■ {data['punchline']}")

    dates = [item["date"] for item in (data["product"] + data["lane1"] + data["lane2"]
                                       + data["direct"] + data["lane3"]) if is_number(item["date"])]
    scale = create_timeline_scale(dates, warnings)
    if scale["type"] == "fitted":
        scale["header_month_keys"] = select_header_month_keys(scale, dates)
        scale["date_center_overrides"] = build_date_center_overrides(dates, scale)

    lower_offset = update_product_row_height(table, data["product"])
    direct_shrink = fit_direct_row(table)

    update_lane(resolver, "product", data["product"], scale, warnings)
    update_lane(resolver, "lane1", data["lane1"], scale, warnings)
    update_lane(resolver, "lane2", data["lane2"], scale, warnings)
    shift_template_lane_vertical(resolver, "lane1", lower_offset)
    shift_template_lane_vertical(resolver, "lane2", lower_offset)
    add_generated_lane(slide, resolver, "direct", data["direct"], scale, warnings, lower_offset)
    add_generated_lane(slide, resolver, "lane3", data["lane3"], scale, warnings, lower_offset - direct_shrink)
    update_subcolumn_labels(table, data)
    update_timeline_header(table, scale, dates)
    move_now_marker(resolver, scale)

    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))

    return {
        "output": str(out_path),
        "warnings": warnings,
        "counts": {
            "product": len(data["product"]),
            "lane1": len(data["lane1"]),
            "lane2": len(data["lane2"]),
            "directSubcolumns": len(data["direct_labels"]),
            "direct": len(data["direct"]),
            "praSubcolumns": len(data["pra_labels"]),
            "lane3": len(data["lane3"]),
        },
    }


def main(argv=None):
    """Thin CLI wrapper around generate_schedule() -- prints the JSON summary."""
    parser = argparse.ArgumentParser(description="Generate schedule PPTX from xlsx (python-pptx port).")
    parser.add_argument("--input", default="input.xlsx")
    parser.add_argument("--template", default=resource_path(os.path.join("assets", "output.pptx")))
    parser.add_argument("--output", default=str(Path("outputs") / "generated_from_input.pptx"))
    args = parser.parse_args(argv)

    summary = generate_schedule(args.input, args.template, args.output)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
